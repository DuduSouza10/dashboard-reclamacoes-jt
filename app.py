from __future__ import annotations

import csv
import io
import os
import re
from datetime import date, datetime
from functools import wraps
from pathlib import Path

from flask import Flask, jsonify, redirect, render_template, request, session, url_for, Response
from openpyxl import load_workbook
from sqlalchemy import Date, Float, Integer, String, create_engine, delete, func, select
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, scoped_session, sessionmaker

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)
SEED_FILE = BASE_DIR / "seed" / "initial_data.csv"


def normalize_database_url(url: str) -> str:
    if url.startswith("postgres://"):
        url = "postgresql://" + url[len("postgres://"):]
    if url.startswith("postgresql://") and "+psycopg" not in url:
        url = "postgresql+psycopg://" + url[len("postgresql://"):]
    return url


DATABASE_URL = normalize_database_url(os.getenv("DATABASE_URL", f"sqlite:///{DATA_DIR / 'dashboard.db'}"))
connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
engine = create_engine(DATABASE_URL, pool_pre_ping=True, future=True, connect_args=connect_args)
DBSession = scoped_session(sessionmaker(bind=engine, autoflush=False, expire_on_commit=False))


class Base(DeclarativeBase):
    pass


class Complaint(Base):
    __tablename__ = "complaints"

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    data: Mapped[date] = mapped_column(Date, index=True)
    regional: Mapped[str] = mapped_column(String(50), default="", index=True)
    numero_ticket: Mapped[str] = mapped_column(String(80), default="", index=True)
    remessa: Mapped[str] = mapped_column(String(80), default="", index=True)
    cliente: Mapped[str] = mapped_column(String(200), default="", index=True)
    base: Mapped[str] = mapped_column(String(120), default="", index=True)
    supervisor: Mapped[str] = mapped_column(String(180), default="", index=True)
    rm: Mapped[str] = mapped_column(String(180), default="", index=True)
    dupla: Mapped[str] = mapped_column(String(180), default="", index=True)
    retorno: Mapped[str] = mapped_column(String(180), default="", index=True)
    valor: Mapped[float] = mapped_column(Float, default=0)
    anexo: Mapped[str] = mapped_column(String(20), default="Não", index=True)
    status: Mapped[str] = mapped_column(String(80), default="Pendente", index=True)


app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "jt-dashboard-reclamacoes-dev-change-me")
app.config["MAX_CONTENT_LENGTH"] = 35 * 1024 * 1024
ASSISTANT_PASSWORD = os.getenv("ASSISTANT_PASSWORD", "3264542")


PT_ZH_RETURNS = {
    "Sem ativo": "未响应",
    "Recebi o produto": "已收到商品",
    "Não recebi o produto": "未收到商品",
    "Já Recebi o Reembolso": "已收到退款",
    "Faltou Produtos": "商品缺失",
    "Produto Diferente": "商品不符",
    "Produto com Defeito": "商品有缺陷",
    "Produto Danificado": "商品损坏",
}


def clean_text(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return re.sub(r"\s+", " ", s)


def parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = clean_text(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Data inválida: {s}")


def parse_float(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    s = clean_text(v).replace("R$", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(".", "").replace(",", ".")
    return float(s)


def row_to_dict(c: Complaint):
    return {
        "id": c.id,
        "data": c.data.isoformat() if c.data else "",
        "regional": c.regional,
        "numero_ticket": c.numero_ticket,
        "remessa": c.remessa,
        "cliente": c.cliente,
        "base": c.base,
        "supervisor": c.supervisor,
        "rm": c.rm,
        "dupla": c.dupla,
        "retorno": c.retorno,
        "valor": round(float(c.valor or 0), 2),
        "anexo": c.anexo,
        "status": c.status,
    }


def normalize_record(record: dict) -> dict:
    aliases = {
        "data": "data",
        "regional": "regional",
        "número do ticket": "numero_ticket",
        "numero do ticket": "numero_ticket",
        "numero_ticket": "numero_ticket",
        "ticket": "numero_ticket",
        "remessa": "remessa",
        "cliente": "cliente",
        "base": "base",
        "supervisor": "supervisor",
        "rm": "rm",
        "dupla": "dupla",
        "retorno": "retorno",
        "valor": "valor",
        "anexo": "anexo",
        "status": "status",
    }
    out = {}
    for key, val in record.items():
        k = clean_text(key).lower()
        if k in aliases:
            out[aliases[k]] = val
    required = ["data", "regional", "numero_ticket", "remessa", "cliente", "base", "supervisor", "rm", "dupla", "retorno", "valor", "anexo", "status"]
    missing = [k for k in required if k not in out]
    if missing:
        raise ValueError("Colunas ausentes: " + ", ".join(missing))
    out["data"] = parse_date(out["data"])
    out["valor"] = parse_float(out["valor"])
    for k in required:
        if k not in ("data", "valor"):
            out[k] = clean_text(out[k])
    out["anexo"] = out["anexo"] or "Não"
    out["status"] = out["status"] or "Pendente"
    return out


def seed_database():
    Base.metadata.create_all(engine)
    db = DBSession()
    try:
        count = db.scalar(select(func.count(Complaint.id))) or 0
        if count == 0 and SEED_FILE.exists():
            batch = []
            with SEED_FILE.open("r", encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f)
                for raw in reader:
                    try:
                        data = normalize_record(raw)
                    except Exception:
                        continue
                    batch.append(Complaint(**data))
                    if len(batch) >= 1000:
                        db.add_all(batch)
                        db.commit()
                        batch.clear()
                if batch:
                    db.add_all(batch)
                    db.commit()
    finally:
        db.close()


seed_database()


def require_assistant(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("assistant_ok"):
            if request.path.startswith("/api/"):
                return jsonify({"error": "Não autorizado"}), 401
            return redirect(url_for("login", next=request.path))
        return fn(*args, **kwargs)
    return wrapper


def filtered_query(args):
    q = select(Complaint)
    if args.get("from"):
        q = q.where(Complaint.data >= parse_date(args["from"]))
    if args.get("to"):
        q = q.where(Complaint.data <= parse_date(args["to"]))
    fields = {
        "regional": Complaint.regional,
        "cliente": Complaint.cliente,
        "base": Complaint.base,
        "supervisor": Complaint.supervisor,
        "rm": Complaint.rm,
        "dupla": Complaint.dupla,
        "retorno": Complaint.retorno,
        "anexo": Complaint.anexo,
        "status": Complaint.status,
    }
    for key, col in fields.items():
        val = clean_text(args.get(key))
        if val and val.lower() not in ("todos", "all"):
            q = q.where(col == val)
    search = clean_text(args.get("q"))
    if search:
        like = f"%{search}%"
        q = q.where(
            Complaint.numero_ticket.ilike(like)
            | Complaint.remessa.ilike(like)
            | Complaint.cliente.ilike(like)
            | Complaint.base.ilike(like)
        )
    return q


@app.teardown_appcontext
def cleanup_db(_exc=None):
    DBSession.remove()


@app.get("/health")
def health():
    return {"ok": True}


@app.get("/")
def dashboard():
    return render_template("dashboard.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        if request.form.get("password", "") == ASSISTANT_PASSWORD:
            session["assistant_ok"] = True
            return redirect(request.args.get("next") or url_for("assistant_panel"))
        error = "Senha incorreta."
    return render_template("login.html", error=error)


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("dashboard"))


@app.get("/assistentes")
@require_assistant
def assistant_panel():
    return render_template("admin.html")


@app.get("/api/filters")
def api_filters():
    db = DBSession()
    try:
        fields = {
            "regional": Complaint.regional,
            "cliente": Complaint.cliente,
            "base": Complaint.base,
            "supervisor": Complaint.supervisor,
            "rm": Complaint.rm,
            "dupla": Complaint.dupla,
            "retorno": Complaint.retorno,
            "anexo": Complaint.anexo,
            "status": Complaint.status,
        }
        out = {}
        for key, col in fields.items():
            vals = db.scalars(select(col).where(col != "").distinct().order_by(col)).all()
            out[key] = vals
        mn, mx = db.execute(select(func.min(Complaint.data), func.max(Complaint.data))).one()
        out["date_min"] = mn.isoformat() if mn else ""
        out["date_max"] = mx.isoformat() if mx else ""
        out["retorno_zh"] = PT_ZH_RETURNS
        return jsonify(out)
    finally:
        db.close()


@app.get("/api/dashboard")
def api_dashboard():
    db = DBSession()
    try:
        rows = db.scalars(filtered_query(request.args)).all()
        total = len(rows)
        total_value = sum(float(r.valor or 0) for r in rows)
        status_counts = {}
        daily = {}
        bases = {}
        returns = {}
        rms = {}
        duplas = {}
        for r in rows:
            st = clean_text(r.status) or "Pendente"
            status_counts[st] = status_counts.get(st, 0) + 1
            day = r.data.isoformat()
            daily.setdefault(day, {})[st] = daily.setdefault(day, {}).get(st, 0) + 1
            bases[r.base] = bases.get(r.base, 0) + 1
            returns[r.retorno] = returns.get(r.retorno, 0) + 1
            for container, key in ((rms, r.rm), (duplas, r.dupla)):
                item = container.setdefault(key, {"name": key, "pedidos": 0, "valor": 0.0, "status": {}, "retornos": {}})
                item["pedidos"] += 1
                item["valor"] += float(r.valor or 0)
                item["status"][st] = item["status"].get(st, 0) + 1
                item["retornos"][r.retorno] = item["retornos"].get(r.retorno, 0) + 1
        delivered = status_counts.get("Entregue", 0)
        pending = status_counts.get("Pendente", 0)
        loss = status_counts.get("Extravio", 0)
        others = total - delivered - pending - loss
        kpis = {
            "pedidos": total,
            "valor": round(total_value, 2),
            "entregue": delivered,
            "pendente": pending,
            "extravio": loss,
            "outros": others,
            "reversao": round(delivered / total, 6) if total else 0,
        }
        daily_list = [{"date": d, **vals} for d, vals in sorted(daily.items())]
        top_bases = sorted(({"name": k, "value": v} for k, v in bases.items() if k), key=lambda x: (-x["value"], x["name"]))[:10]
        top_returns = sorted(({"name": k or "Sem retorno", "value": v} for k, v in returns.items()), key=lambda x: -x["value"])
        rm_rows = sorted(rms.values(), key=lambda x: (-x["pedidos"], x["name"]))
        dupla_rows = sorted(duplas.values(), key=lambda x: (-x["pedidos"], x["name"]))
        return jsonify({
            "kpis": kpis,
            "status_counts": status_counts,
            "daily": daily_list,
            "top_bases": top_bases,
            "returns": top_returns,
            "rm_rows": rm_rows,
            "dupla_rows": dupla_rows,
        })
    finally:
        db.close()


@app.get("/api/records")
def api_records():
    db = DBSession()
    try:
        page = max(int(request.args.get("page", 1)), 1)
        per_page = min(max(int(request.args.get("per_page", 50)), 10), 200)
        base_q = filtered_query(request.args)
        count_q = select(func.count()).select_from(base_q.subquery())
        total = db.scalar(count_q) or 0
        q = base_q.order_by(Complaint.data.desc(), Complaint.id.desc()).offset((page - 1) * per_page).limit(per_page)
        rows = db.scalars(q).all()
        return jsonify({"rows": [row_to_dict(r) for r in rows], "total": total, "page": page, "per_page": per_page})
    finally:
        db.close()


@app.post("/api/record")
@require_assistant
def api_add_record():
    db = DBSession()
    try:
        data = normalize_record(request.get_json(force=True))
        obj = Complaint(**data)
        db.add(obj)
        db.commit()
        return jsonify(row_to_dict(obj)), 201
    except Exception as e:
        db.rollback()
        return jsonify({"error": str(e)}), 400
    finally:
        db.close()


@app.get("/api/record/<int:record_id>")
@require_assistant
def api_get_record(record_id: int):
    db = DBSession()
    try:
        obj = db.get(Complaint, record_id)
        if not obj:
            return jsonify({"error": "Registro não encontrado"}), 404
        return jsonify(row_to_dict(obj))
    finally:
        db.close()


@app.patch("/api/record/<int:record_id>")
@require_assistant
def api_update_record(record_id: int):
    db = DBSession()
    try:
        obj = db.get(Complaint, record_id)
        if not obj:
            return jsonify({"error": "Registro não encontrado"}), 404
        payload = request.get_json(force=True)
        allowed = {"data", "regional", "numero_ticket", "remessa", "cliente", "base", "supervisor", "rm", "dupla", "retorno", "valor", "anexo", "status"}
        for k, v in payload.items():
            if k not in allowed:
                continue
            if k == "data":
                v = parse_date(v)
            elif k == "valor":
                v = parse_float(v)
            else:
                v = clean_text(v)
            setattr(obj, k, v)
        db.commit()
        return jsonify(row_to_dict(obj))
    except Exception as e:
        db.rollback()
        return jsonify({"error": str(e)}), 400
    finally:
        db.close()


@app.delete("/api/record/<int:record_id>")
@require_assistant
def api_delete_record(record_id: int):
    db = DBSession()
    try:
        obj = db.get(Complaint, record_id)
        if not obj:
            return jsonify({"error": "Registro não encontrado"}), 404
        db.delete(obj)
        db.commit()
        return jsonify({"ok": True})
    finally:
        db.close()


@app.post("/api/bulk-update")
@require_assistant
def api_bulk_update():
    db = DBSession()
    try:
        payload = request.get_json(force=True)
        ids = [int(x) for x in payload.get("ids", [])]
        if not ids:
            return jsonify({"error": "Nenhum registro selecionado"}), 400
        values = {k: clean_text(v) for k, v in payload.get("values", {}).items() if k in {"status", "anexo", "retorno"} and clean_text(v)}
        if not values:
            return jsonify({"error": "Nenhuma alteração informada"}), 400
        rows = db.scalars(select(Complaint).where(Complaint.id.in_(ids))).all()
        for row in rows:
            for k, v in values.items():
                setattr(row, k, v)
        db.commit()
        return jsonify({"ok": True, "updated": len(rows)})
    except Exception as e:
        db.rollback()
        return jsonify({"error": str(e)}), 400
    finally:
        db.close()


def records_from_upload(file_storage):
    name = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        return list(csv.DictReader(io.StringIO(text), delimiter=delimiter))
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True, keep_links=False)
        preferred = next((n for n in wb.sheetnames if clean_text(n).lower() == "status pedido"), wb.sheetnames[0])
        ws = wb[preferred]
        rows = ws.iter_rows(values_only=True)
        headers = [clean_text(x) for x in next(rows)]
        out = []
        for values in rows:
            if not any(v is not None and clean_text(v) for v in values):
                continue
            out.append(dict(zip(headers, values)))
        return out
    raise ValueError("Use arquivo .xlsx ou .csv")


@app.post("/api/import")
@require_assistant
def api_import():
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "Selecione um arquivo"}), 400
    mode = request.form.get("mode", "merge")
    db = DBSession()
    try:
        raw_records = records_from_upload(f)
        normalized = []
        errors = []
        for i, raw in enumerate(raw_records, start=2):
            try:
                normalized.append(normalize_record(raw))
            except Exception as e:
                if len(errors) < 20:
                    errors.append(f"Linha {i}: {e}")
        if mode == "replace":
            db.execute(delete(Complaint))
            db.commit()
            for start in range(0, len(normalized), 1000):
                db.add_all([Complaint(**r) for r in normalized[start:start+1000]])
                db.commit()
            return jsonify({"ok": True, "mode": mode, "imported": len(normalized), "errors": errors})

        existing_rows = db.scalars(select(Complaint)).all()
        index = {(clean_text(r.numero_ticket), clean_text(r.remessa)): r for r in existing_rows}
        inserted = updated = 0
        for rec in normalized:
            key = (clean_text(rec["numero_ticket"]), clean_text(rec["remessa"]))
            obj = index.get(key)
            if obj:
                for k, v in rec.items():
                    setattr(obj, k, v)
                updated += 1
            else:
                obj = Complaint(**rec)
                db.add(obj)
                index[key] = obj
                inserted += 1
        db.commit()
        return jsonify({"ok": True, "mode": mode, "inserted": inserted, "updated": updated, "errors": errors})
    except Exception as e:
        db.rollback()
        return jsonify({"error": str(e)}), 400
    finally:
        db.close()


@app.get("/export.csv")
@require_assistant
def export_csv():
    db = DBSession()
    try:
        rows = db.scalars(filtered_query(request.args).order_by(Complaint.data.desc(), Complaint.id.desc())).all()
        out = io.StringIO()
        writer = csv.writer(out, delimiter=';')
        writer.writerow(["Data", "Regional", "Número do ticket", "Remessa", "Cliente", "Base", "Supervisor", "RM", "Dupla", "Retorno", "Valor", "Anexo", "Status"])
        for r in rows:
            writer.writerow([r.data.strftime("%d/%m/%Y"), r.regional, r.numero_ticket, r.remessa, r.cliente, r.base, r.supervisor, r.rm, r.dupla, r.retorno, f"{float(r.valor or 0):.2f}".replace('.', ','), r.anexo, r.status])
        payload = '\ufeff' + out.getvalue()
        return Response(payload, mimetype="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=status_pedido.csv"})
    finally:
        db.close()


if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=os.getenv("FLASK_DEBUG") == "1")
